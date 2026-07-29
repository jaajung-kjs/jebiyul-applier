"""적용근거 시트를 코드로 렌더링해 결과 xlsx를 만든다(템플릿 복사 없음).

입력(params)+적용율(rates)을 sheet_model이 블록 리스트로 바꾸고, renderer가 그것을
빈 워크북에 그려 저장한다. 양식 자체가 코드로 정의되므로 템플릿 파일이 필요 없다.
"""
from openpyxl import Workbook

from src import lookup  # noqa: F401  (호환 import 유지)
from src import sheet_model
from src import nomu_model
from src import renderer
from src import styles
from src.hwp_reader import read_hwp

# Canonical item list — single source of truth for the runtime package.
ITEMS = [
    "간접노무비",
    "공구손료",
    "산재보험료",
    "고용보험료",
    "건강보험료",
    "연금보험료",
    "퇴직공제부금비",
    "노인장기요양보험료",
    "산업안전보건관리비",
    "기타경비",
    "일반관리비",
    "이윤",
]


def build_output(rates: dict, out_path: str, params: dict | None = None,
                 jebiyul_path: str | None = None, template_path=None,
                 hwp_path: str | None = None, selected_nomu=None) -> str:
    """rates로 적용근거를, (hwp_path+selected_nomu가 있으면) 노무임 시트까지 생성.

    hwp 인자가 없으면 기존과 동일하게 '적용근거' 단일 시트만 만든다.

    Args:
        rates:         compute_rates 결과(12개 항목 → 적용율).
        out_path:      저장 경로.
        params:        공사 파라미터. None이면 구간표 없이 골격만 그린다.
        jebiyul_path:  제비율 파일 경로. params와 함께 있어야 구간표를 채운다.
        template_path: (deprecated) 과거 템플릿 복사 방식 호환용 인자, 무시된다.
        hwp_path:      임금실태조사 hwp 경로. selected_nomu와 함께 있어야
                       '7.통신노무임' 시트가 추가된다.
        selected_nomu: nomu_model.build에 넘길 선택 직종명 리스트.

    Returns:
        저장된 파일 경로.
    """
    blocks8 = sheet_model.build(params or _params_from_rates(), rates, jebiyul_path)
    if not (hwp_path and selected_nomu):
        return renderer.render(blocks8, out_path)

    report = read_hwp(hwp_path)
    blocks7 = nomu_model.build(report, selected_nomu)
    wb = Workbook()
    wb.remove(wb.active)
    renderer.render_sheet(wb.create_sheet("7.통신노무임"), blocks7,
                          styles.NOMU_COL_WIDTHS)
    renderer.render_sheet(wb.create_sheet("8.적용근거"), blocks8)
    wb.save(out_path)
    return out_path


def _params_from_rates() -> dict:
    """params 없이 호출되는 골격 경로용 최소 파라미터(구간표 생략)."""
    return {"kind": "", "jikjeop_cost": 0, "days": 0,
            "contract": "경쟁", "sanjae_basis": "한전", "sanan_target": 0}
