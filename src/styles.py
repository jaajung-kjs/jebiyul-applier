"""PIU 적용근거 시트에서 추출한 스타일 팔레트. 렌더러가 셀에 입히는 폰트·채움·테두리·정렬·치수."""
from openpyxl.styles import Font, PatternFill, Border, Side, Alignment

BODY_FONT_NAME = "맑은 고딕"
TITLE_FONT_NAME = "HY견고딕"

HEADER_FILL_RGB = "FFDBE5F1"      # accent1 80% lighter
HIGHLIGHT_FILL_RGB = "FFFFF2CC"   # 적용행 강조(옅은 금색)
SECTION_COLOR = "FF0000FF"        # 섹션 제목(파랑)
RATE_COLOR = "FFFF0000"           # 적용율(빨강)

COL_WIDTHS = {"A": 4.2, "B": 10.2, "C": 14.1, "D": 10.2, "E": 12.6,
              "G": 10.2, "H": 12.4, "I": 10.2, "K": 2.5}
# 직종명(A) · No.(B) · 공표일 과거(C·D·E) · 현재(F) · 변동율(G) · 비고(H)
NOMU_COL_WIDTHS = {"A": 16.0, "B": 7.0, "C": 11.0, "D": 11.0, "E": 11.0,
                   "F": 11.0, "G": 9.0, "H": 12.0}
COMMA_FMT = "#,##0"
PCT_FMT = "0.0%"
ROW_TITLE_H = 35.1
ROW_BODY_H = 20.1


def body_font():
    return Font(name=BODY_FONT_NAME, size=11)


def title_font():
    return Font(name=TITLE_FONT_NAME, size=24)


def bold_font():
    return Font(name=BODY_FONT_NAME, size=11, bold=True)


def section_font():
    return Font(name=BODY_FONT_NAME, size=11, bold=True, color=SECTION_COLOR)


def rate_font():
    return Font(name=BODY_FONT_NAME, size=11, bold=True, color=RATE_COLOR)


def header_fill():
    return PatternFill(patternType="solid", fgColor=HEADER_FILL_RGB)


def highlight_fill():
    return PatternFill(patternType="solid", fgColor=HIGHLIGHT_FILL_RGB)


def side(style):
    return Side(style=style)


def box_border(outer="thin", inner="hair"):
    return Border(top=side(outer), bottom=side(outer),
                  left=side(outer), right=side(outer))


def center(wrap=False):
    return Alignment(horizontal="center", vertical="center", wrap_text=wrap)


def left():
    return Alignment(horizontal="left", vertical="center")


def right():
    return Alignment(horizontal="right", vertical="center")
