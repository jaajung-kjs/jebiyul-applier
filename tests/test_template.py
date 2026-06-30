"""Tests for the standard 적용근거 template (Task 7).

Pins the row-layout contract so Task 8's builder can rely on it.
"""
import os
import openpyxl
import pytest

from src.builder import ITEMS

TEMPLATE_PATH = os.path.join(
    os.path.dirname(os.path.dirname(__file__)),
    "assets",
    "template_적용근거.xlsx",
)


@pytest.fixture(scope="module")
def ws():
    wb = openpyxl.load_workbook(TEMPLATE_PATH)
    return wb["적용근거"]


def test_template_file_exists():
    assert os.path.isfile(TEMPLATE_PATH), f"Template not found: {TEMPLATE_PATH}"


def test_sheet_name(ws):
    # fixture already selected the sheet by name; if it raised KeyError the test fails
    assert ws.title == "적용근거"


def test_template_has_all_items(ws):
    """All 12 item labels must appear somewhere in column A."""
    col_a = {ws.cell(row=r, column=1).value for r in range(1, ws.max_row + 1)}
    for item in ITEMS:
        assert item in col_a, f"Label '{item}' missing from column A"


def test_item_row_positions(ws):
    """Each label must sit at row = 3 + i*2 (Task 8 CELL_MAP contract)."""
    for i, item in enumerate(ITEMS):
        expected_row = 3 + i * 2
        actual = ws.cell(row=expected_row, column=1).value
        assert actual == item, (
            f"Item {i} ('{item}') expected at row {expected_row}, "
            f"got '{actual}'"
        )


def test_rate_cells_empty(ws):
    """Column I at each item row must be empty (builder fills them)."""
    for i in range(len(ITEMS)):
        row = 3 + i * 2
        val = ws.cell(row=row, column=9).value
        assert val is None, (
            f"I{row} should be empty but contains '{val}'"
        )
