import openpyxl
from openpyxl import Workbook
from src import sheet_model as M
from src import renderer


def _blocks():
    return [
        M.Title("공사비 산출 적용근거"),
        M.SectionHeader("1. 간접노무비", M.NOTE),
        M.NoteLines(["    ☞ 계상금액 : 직접노무비 × 적용율"]),
        M.AppliedRate("    ☞ 적 용 율 :  ", 0.126, fmt="0.0%"),
    ]


def test_render_sheet_draws_into_given_ws(tmp_path):
    from src import renderer, styles
    wb = Workbook()
    ws = wb.active
    ws.title = "직접"
    renderer.render_sheet(ws, _blocks(), col_widths=styles.COL_WIDTHS)
    assert ws["A2"].value == "공사비 산출 적용근거"


def test_render_writes_title_merged(tmp_path):
    out = str(tmp_path / "o.xlsx")
    renderer.render(_blocks(), out)
    ws = openpyxl.load_workbook(out)["적용근거"]
    assert ws["A2"].value == "공사비 산출 적용근거"
    assert "A2:J2" in [str(m) for m in ws.merged_cells.ranges]


def test_section_header_blue_bold(tmp_path):
    out = str(tmp_path / "o.xlsx")
    renderer.render(_blocks(), out)
    ws = openpyxl.load_workbook(out)["적용근거"]
    found = None
    for row in ws.iter_rows():
        for c in row:
            if c.value == "1. 간접노무비":
                found = c
    assert found is not None
    assert found.font.bold is True
    assert found.font.color.rgb == "FF0000FF"


def test_applied_rate_value_in_I_with_percent_format(tmp_path):
    out = str(tmp_path / "o.xlsx")
    renderer.render(_blocks(), out)
    ws = openpyxl.load_workbook(out)["적용근거"]
    hit = None
    for c in ws["I"]:
        if isinstance(c.value, (int, float)) and abs(c.value - 0.126) < 1e-9:
            hit = c
    assert hit is not None
    assert hit.number_format == "0.0%"
    assert hit.font.color.rgb == "FFFF0000"


def test_highlight_skips_vertically_merged_label(tmp_path):
    """세로 병합된 규모 라벨 칸은 강조색을 받지 않고, 단일 행 셀만 강조된다."""
    t = M.BandTable(
        kind="gibon",
        headers=[("공사규모", 2, 2), ("적용율", 5, 6)],
        rows=[
            M.BandRow([("50억 미만", 2, 2, None, 2), ("183일", 3, 4, None), (0.19, 5, 6, "0.0%")],
                      highlight=True),
            M.BandRow([("365일", 3, 4, None), (0.20, 5, 6, "0.0%")]),
        ],
        criterion="직접노무비 × 적용율",
    )
    out = str(tmp_path / "m.xlsx")
    renderer.render([t], out)
    ws = openpyxl.load_workbook(out)["적용근거"]
    size_cell = None
    dur_cell = None
    for row in ws.iter_rows():
        for c in row:
            if c.value == "50억 미만":
                size_cell = c
            if c.value == "183일":
                dur_cell = c
    assert size_cell.fill.patternType != "solid"      # 병합 라벨은 강조 안 됨
    assert dur_cell.fill.fgColor.rgb == "FFFFF2CC"     # 단일 행 셀은 강조


def test_applied_rate_annotation_in_J(tmp_path):
    blocks = [M.AppliedRate("    ☞ 적 용 율 :  (사급재료비 제외시)", 0.0315,
                            fmt="0.000%", annotation="× 1.2")]
    out = str(tmp_path / "a.xlsx")
    renderer.render(blocks, out)
    ws = openpyxl.load_workbook(out)["적용근거"]
    assert any(c.value == "× 1.2" for c in ws["J"])


def test_column_widths_applied(tmp_path):
    out = str(tmp_path / "o.xlsx")
    renderer.render(_blocks(), out)
    ws = openpyxl.load_workbook(out)["적용근거"]
    assert round(ws.column_dimensions["A"].width, 1) == 4.2
    assert round(ws.column_dimensions["K"].width, 1) == 2.5


def test_table_header_filled_and_highlight(tmp_path):
    t = M.BandTable(
        kind="ilban",
        headers=[("공 사 규 모", 2, 3), ("적용율", 4, 6), ("적  용  기  준", 7, 10)],
        rows=[
            M.BandRow([("5억원 미만", 2, 3, None), (0.06, 4, 6, "0.0%")], highlight=True),
            M.BandRow([("100억원 이상", 2, 3, None), (0.045, 4, 6, "0.0%")]),
        ],
        criterion="[지입재료비+노무비+도급분경비]×적용율",
    )
    out = str(tmp_path / "t.xlsx")
    renderer.render([t], out)
    ws = openpyxl.load_workbook(out)["적용근거"]
    hdr = None
    hirow = None
    for row in ws.iter_rows():
        for c in row:
            if c.value == "공 사 규 모":
                hdr = c
            if c.value == "5억원 미만":
                hirow = c
    assert hdr.fill.fgColor.rgb == "FFDBE5F1"
    assert hirow.fill.fgColor.rgb == "FFFFF2CC"


def _nomu_ws(N, renderer, styles):
    blocks = [
        N.NomuHeader(["2024.9.1", "2025.1.1", "2025.9.1"], "2026.1.1"),
        N.NomuGroup("Ⅰ", "일반공사직종"),
        N.NomuRow(75, "보통인부", "1002", [167081, 169804, 171037], 172068, 0.006028, ""),
        N.NomuAvg(0.006028),
    ]
    wb = __import__("openpyxl").Workbook()
    ws = wb.active
    renderer.render_sheet(ws, blocks, styles.NOMU_COL_WIDTHS)
    return ws


def test_nomu_title_matches_apply_sheet_style(tmp_path):
    """노무임 제목이 적용근거 제목과 같은 title_font로 크게 렌더된다."""
    from src import renderer, styles, nomu_model as N
    wb = __import__("openpyxl").Workbook()
    ws = wb.active
    renderer.render_sheet(ws, [N.NomuTitle("시중노무임 산출")], styles.NOMU_COL_WIDTHS)
    c = next(c for row in ws.iter_rows() for c in row if c.value == "시중노무임 산출")
    assert c.font.name == styles.TITLE_FONT_NAME
    assert c.font.size == 24


def test_nomu_render_row_and_formulas(tmp_path):
    from src import renderer, styles, nomu_model as N
    ws = _nomu_ws(N, renderer, styles)
    vals = [c.value for row in ws.iter_rows() for c in row if c.value not in (None, "")]
    assert "보통인부" in vals
    assert 172068 in vals            # 현재 노임(실제 적용값)
    assert "노임변동률평균" in vals
    # 현재 노임 셀은 콤마 서식
    cur = next(c for row in ws.iter_rows() for c in row if c.value == 172068)
    assert cur.number_format == "#,##0"
    r = cur.row  # 보통인부 데이터 행
    # 변동율은 하드코딩 값이 아니라 계산식: =(현재-직전)/직전
    delta = ws.cell(r, 8)   # H열(과거3 → 변동율)
    assert delta.value == "=(G%d-F%d)/F%d" % (r, r, r)
    assert delta.number_format == "0.0%"
    # 노임변동률평균도 계산식(AVERAGE)
    avg_r = next(c.row for row in ws.iter_rows() for c in row if c.value == "노임변동률평균")
    avg = ws.cell(avg_r, 8)
    assert avg.value == "=AVERAGE(H%d:H%d)" % (r, r)
    assert avg.number_format == "0.0%"


def test_nomu_current_column_highlighted(tmp_path):
    """실제 적용되는 현재 시점 열이 표 전체 높이에 걸쳐 끊김 없이 강조된다."""
    from src import renderer, styles, nomu_model as N
    ws = _nomu_ws(N, renderer, styles)
    cur = next(c for row in ws.iter_rows() for c in row if c.value == 172068)
    cur_c = cur.column  # 현재 노임(=적용) 열
    # 헤더·데이터·그룹헤더·평균 행 모두 현재열이 강조된다(연속 띠)
    labels = ("2026.1.1", 172068, "Ⅰ. 일반공사직종", "노임변동률평균")
    for label in labels:
        r = next(c.row for row in ws.iter_rows() for c in row if c.value == label)
        cell = ws.cell(r, cur_c)
        assert cell.fill.fgColor.rgb == styles.HIGHLIGHT_FILL_RGB, f"{label} 현재열 강조 없음"
    # 본문(그룹헤더~평균) 현재열에 강조 공백이 없다(띄엄띄엄 금지).
    # (헤더 현재열은 rowspan 병합이라 구성원 셀 fill은 openpyxl로 안 읽혀 범위서 제외)
    top = next(c.row for row in ws.iter_rows() for c in row if c.value == "Ⅰ. 일반공사직종")
    bot = max(c.row for row in ws.iter_rows() for c in row if c.value == "노임변동률평균")
    for r in range(top, bot + 1):
        cell = ws.cell(r, cur_c)
        assert cell.fill.fgColor.rgb == styles.HIGHLIGHT_FILL_RGB, f"r{r} 현재열 강조 끊김"


def test_nomu_all_cells_have_full_borders(tmp_path):
    """모든 행(그룹 헤더·평균 포함)의 모든 칸에 4면 테두리가 들어간다."""
    from src import renderer, styles, nomu_model as N
    ws = _nomu_ws(N, renderer, styles)

    def full(c):
        b = c.border
        return all(getattr(b, k) is not None and getattr(b, k).style
                   for k in ("top", "bottom", "left", "right"))

    note_c = 9  # 과거 3열 → 비고 = I열
    for label in ("Ⅰ. 일반공사직종", "노임변동률평균", "보통인부"):
        r = next(c.row for row in ws.iter_rows() for c in row if c.value == label)
        for col in range(1, note_c + 1):
            assert full(ws.cell(r, col)), f"{label} r{r} c{col} 테두리 누락"
