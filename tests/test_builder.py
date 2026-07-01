"""src/builder.py — model→render 파이프라인 검증(템플릿 복사 폐지)."""
import os

import openpyxl

from src.builder import build_output, ITEMS

SAMPLE_RATES = {
    "간접노무비": 0.195,
    "공구손료": 0.03,
    "산재보험료": 0.03656,
    "고용보험료": 0.0101,
    "건강보험료": 0.03595,
    "연금보험료": 0.0475,
    "퇴직공제부금비": 0.023,
    "노인장기요양보험료": 0.1314,
    "산업안전보건관리비": 0.0315,
    "기타경비": 0.057,
    "일반관리비": 0.08,
    "이윤": 0.15,
}


def _texts(ws):
    return [ws.cell(r, c).value
            for r in range(1, ws.max_row + 1)
            for c in range(1, ws.max_column + 1)]


def test_build_returns_path(tmp_path):
    out_path = str(tmp_path / "out.xlsx")
    assert build_output(SAMPLE_RATES, out_path) == out_path
    assert os.path.isfile(out_path)


def test_skeleton_has_title_and_sections(tmp_path):
    """params 없이 호출하면 표 없는 골격을 그린다(제목·섹션 존재)."""
    out_path = str(tmp_path / "out.xlsx")
    build_output(SAMPLE_RATES, out_path)
    ws = openpyxl.load_workbook(out_path)["적용근거"]
    assert ws["A2"].value == "공사비 산출 적용근거"
    texts = _texts(ws)
    assert "1. 간접노무비" in texts
    assert any(str(t).startswith("4. 이") for t in texts if t)


def test_skeleton_omits_band_tables(tmp_path):
    """jebiyul_path 없으면 구간표 헤더가 그려지지 않는다."""
    out_path = str(tmp_path / "out.xlsx")
    build_output(SAMPLE_RATES, out_path)
    ws = openpyxl.load_workbook(out_path)["적용근거"]
    assert "공사규모" not in _texts(ws)


def test_inline_rate_text_uses_rates(tmp_path):
    """설명 문구에 박힌 보험요율이 rates 값으로 써진다."""
    out = build_output(SAMPLE_RATES, str(tmp_path / "out.xlsx"))
    ws = openpyxl.load_workbook(out)["적용근거"]
    texts = [str(t) for t in _texts(ws) if t]
    assert any("3.656%" in t for t in texts)   # 산재
    assert any("3.595%" in t for t in texts)   # 건강
    assert any("13.14%" in t for t in texts)   # 노인장기요양


def test_tables_filled_from_file(tomok_path, tmp_path):
    """params+제비율 경로가 주어지면 구간표가 파일값으로 채워진다(<50억 압축형)."""
    params = dict(jikjeop_cost=500_000_000, days=200, kind="토목", contract="경쟁",
                  sanjae_basis="한전", sanan_target=400_000_000, jebiyul_path=tomok_path)
    out = build_output(SAMPLE_RATES, str(tmp_path / "out.xlsx"),
                       params=params, jebiyul_path=tomok_path)
    ws = openpyxl.load_workbook(out)["적용근거"]
    nums = [v for v in _texts(ws) if isinstance(v, (int, float))]
    texts = [str(v) for v in _texts(ws) if v]
    # 간접노무비표 50억미만 라벨 + 183일 값(0.191)
    assert any("50억 미만" in t for t in texts)
    assert any(abs(v - 0.191) < 1e-9 for v in nums)
    # 이윤표 경쟁 50억미만 = 0.15
    assert any(abs(v - 0.15) < 1e-9 for v in nums)
    # 산안비 5억미만 구간(0.0315)
    assert any(abs(v - 0.0315) < 1e-9 for v in nums)


def test_items_list_has_twelve():
    assert len(ITEMS) == 12
    assert ITEMS[0] == "간접노무비"
    assert ITEMS[-1] == "이윤"
