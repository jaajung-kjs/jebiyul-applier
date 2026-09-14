"""율-테이블 골든 테스트. 실제 xlsx 파일의 셀 값으로 검증.

파일 기반 골든 값 (probe 확인):
  토목 파일 ('26. 4. 13.' 시트):
    AC15=19.1, AC17=19.5, AF15=18.4   (간접노무비)
    AL15=5.5, AO15=4.9               (기타경비)
    BD15=8,   BD43=4.5               (일반관리비, col56)
    BS15=15,  BS43=9                 (이윤, col71)
    BD15=8 (10-50억 merged → row 15  (일반관리비 10-50억, probe=8)
    BS43=9  (이윤 1000억이상 경쟁 파일값, probe=9)
  건축 파일 ('건축제비율(4.13.)' 시트):
    AC15=17.5                        (간접노무비)
    AL15=5.0                         (기타경비)
    BE15=8                           (일반관리비, col57)
    BR15=15                          (이윤, col70)
"""
import pytest
from src.lookup import table_rate


# ── 간접노무비 (golden: 브리핑 제공 검증값) ──────────────────────────────

def test_tomok_ganjeop_10eok_183(tomok_path):
    assert table_rate(tomok_path, "간접노무비", "토목", "10억미만", "183") == pytest.approx(0.191)

def test_tomok_ganjeop_10eok_365(tomok_path):
    assert table_rate(tomok_path, "간접노무비", "토목", "10억미만", "365") == pytest.approx(0.195)

def test_tomok_ganjeop_jogyeong(tomok_path):
    assert table_rate(tomok_path, "간접노무비", "조경", "10억미만", "183") == pytest.approx(0.184)

def test_geonchuk_ganjeop_10eok_183(geonchuk_path):
    assert table_rate(geonchuk_path, "간접노무비", "건축", "10억미만", "183") == pytest.approx(0.175)


# ── 기타경비 (probe 확인: AL15=5.5, AO15=4.9) ──────────────────────────

def test_tomok_gita_10eok_183(tomok_path):
    assert table_rate(tomok_path, "기타경비", "토목", "10억미만", "183") == pytest.approx(0.055)

def test_tomok_gita_jogyeong_10eok_183(tomok_path):
    assert table_rate(tomok_path, "기타경비", "조경", "10억미만", "183") == pytest.approx(0.049)

def test_geonchuk_gita_10eok_183(geonchuk_path):
    # 건축 파일 AL15=5.0
    assert table_rate(geonchuk_path, "기타경비", "건축", "10억미만", "183") == pytest.approx(0.05)


# ── 일반관리비 (기간 무관; probe 확인: BD15=8, BD43=4.5) ────────────────

def test_tomok_ilban_10eok(tomok_path):
    # 공사원가 5억미만 행(BD15) — duration은 무시됨
    assert table_rate(tomok_path, "일반관리비", "토목", "10억미만", "183") == pytest.approx(0.08)

def test_tomok_ilban_1000eok(tomok_path):
    # 공사원가 1000억이상 행(BD43)
    assert table_rate(tomok_path, "일반관리비", "토목", "1000억이상", "183") == pytest.approx(0.045)

def test_geonchuk_ilban_10eok(geonchuk_path):
    # 건축 파일 BE15=8
    assert table_rate(geonchuk_path, "일반관리비", "건축", "10억미만", "183") == pytest.approx(0.08)


# ── 이윤 (기간 무관; probe 확인: BS15=15, BS43=9) ───────────────────────

def test_tomok_iyun_10eok(tomok_path):
    assert table_rate(tomok_path, "이윤", "토목", "10억미만", "183") == pytest.approx(0.15)

def test_tomok_iyun_1000eok(tomok_path):
    assert table_rate(tomok_path, "이윤", "토목", "1000억이상", "183") == pytest.approx(0.09)

def test_geonchuk_iyun_10eok(geonchuk_path):
    # 건축 파일 BR15=15
    assert table_rate(geonchuk_path, "이윤", "건축", "10억미만", "183") == pytest.approx(0.15)


# ── Fix 1: LookupError 경로 테스트 ──────────────────────────────────────────

def test_bad_label_raises(tomok_path):
    import openpyxl
    from src.lookup import cell_text_search
    ws = openpyxl.load_workbook(tomok_path, data_only=True).worksheets[0]
    with pytest.raises(LookupError):
        cell_text_search(ws, "[절대없는라벨XYZ]")

def test_no_data_sheet_raises(tmp_path):
    import openpyxl
    from src.lookup import find_data_sheet
    wb = openpyxl.Workbook(); wb.active.title = "빈시트"
    with pytest.raises(LookupError):
        find_data_sheet(wb)


# ── Fix 2: 수의 이윤율 하드코딩 ──────────────────────────────────────────────

def test_iyun_gyeongjaeng_uses_file(tomok_path):
    # 경쟁(또는 None) → 파일값. 1000억이상 = 0.09 (probe 확인: BS43=9)
    assert table_rate(tomok_path, "이윤", "토목", "1000억이상", "183", "경쟁") == pytest.approx(0.09)

def test_iyun_suui_under_1000_hardcoded(tomok_path):
    # 수의계약 + 1000억미만 → 조달청 기준 하드코딩 0.10
    assert table_rate(tomok_path, "이윤", "토목", "50-300억", "183", "수의") == pytest.approx(0.10)

def test_iyun_suui_over_1000_hardcoded(tomok_path):
    # 수의계약 + 1000억이상 → 조달청 기준 하드코딩 0.09
    assert table_rate(tomok_path, "이윤", "토목", "1000억이상", "183", "수의") == pytest.approx(0.09)


# ── Fix 3: 병합 셀 워크어라운드 고정 테스트 ──────────────────────────────────

def test_tomok_ilban_10_50eok_merged(tomok_path):
    # 10-50억 → SIZE_BASE_ROW_ILBAN=15 (BD15:BK30 merged, probe 확인: BD15=8)
    assert table_rate(tomok_path, "일반관리비", "토목", "10-50억", "183") == pytest.approx(0.08)


# ── Fix 4: 전기·통신·소방·전문 일반관리비 전용 열 ────────────────────────────
# 제비율 파일 주석: "전기∙통신∙소방∙전문 및 기타공사의 경우 일반관리비요율을 제외한
# 각종 요율은 토목, 건축 등 관련 공사업종에 따라 적용" — 즉 일반관리비만 전문 열.
# 전문 열(토목파일 col64 / 건축파일 col63)은 종합 열과 구간별 요율이 다르다.

def test_jeonmun_ilban_uses_specialty_column_tomok(tomok_path):
    """전문 일반관리비: 종합(8/8/6.5/5/4.5)이 아니라 전문(8/6.5/5/4.5/4.5)."""
    f = lambda s: table_rate(tomok_path, "일반관리비", "전기통신소방전문", s, "183")
    assert f("10억미만") == pytest.approx(0.08)
    assert f("10-50억") == pytest.approx(0.065)      # 종합이면 0.08
    assert f("50-300억") == pytest.approx(0.05)      # 종합이면 0.065
    assert f("300-1000억") == pytest.approx(0.045)   # 종합이면 0.05
    assert f("1000억이상") == pytest.approx(0.045)


def test_jeonmun_ilban_uses_specialty_column_geonchuk(geonchuk_path):
    """건축 파일은 열 위치가 달라도(전문 col63) 헤더로 찾아 동작해야 한다."""
    f = lambda s: table_rate(geonchuk_path, "일반관리비", "전기통신소방전문", s, "183")
    assert f("10-50억") == pytest.approx(0.065)
    assert f("50-300억") == pytest.approx(0.05)


def test_jeonmun_other_items_still_use_general_column(tomok_path):
    """일반관리비 외 항목은 그대로 토목 종합열을 쓴다(파일 주석 규칙)."""
    assert (table_rate(tomok_path, "간접노무비", "전기통신소방전문", "10억미만", "183")
            == table_rate(tomok_path, "간접노무비", "토목", "10억미만", "183"))


def test_general_kinds_ilban_unchanged(tomok_path, geonchuk_path):
    """회귀 방지: 종합 공사종류의 일반관리비 값은 기존과 동일해야 한다."""
    g = lambda s: table_rate(tomok_path, "일반관리비", "토목", s, "183")
    assert [g(s) for s in ("10억미만", "10-50억", "50-300억", "300-1000억", "1000억이상")] == \
           [pytest.approx(x) for x in (0.08, 0.08, 0.065, 0.05, 0.045)]
    assert table_rate(geonchuk_path, "일반관리비", "건축", "10-50억", "183") == pytest.approx(0.08)


def test_iyun_unchanged_after_row_remap(tomok_path):
    """회귀 방지: 이윤도 SIZE_BASE_ROW_ILBAN을 쓰므로 값이 그대로여야 한다."""
    g = lambda s: table_rate(tomok_path, "이윤", "토목", s, "183", "경쟁")
    assert [g(s) for s in ("10억미만", "10-50억", "50-300억", "300-1000억", "1000억이상")] == \
           [pytest.approx(x) for x in (0.15, 0.15, 0.12, 0.10, 0.09)]


# ── Fix 5: 공사종류 열을 헤더 텍스트로 찾는다(열 번호 하드코딩 제거) ──────────
# 토목 파일: 간접노무비 토목=29, 조경=32, 산업설비(토목)=35
# 건축 파일: 간접노무비 건축=29, 산업설비(건축)=34  ← 기존 DATA_COL은 35로 오인
#            (병합 셀에 우연히 걸쳐 값만 맞았을 뿐 열 자체가 틀렸다)

def _col(path, item, kind):
    import openpyxl
    from src.lookup import find_data_sheet, _col_for_item_kind
    ws = find_data_sheet(openpyxl.load_workbook(path, data_only=True))
    return _col_for_item_kind(ws, item, kind, kind)


def test_kind_col_resolved_by_header_tomok(tomok_path):
    assert _col(tomok_path, "간접노무비", "토목") == 29
    assert _col(tomok_path, "간접노무비", "조경") == 32
    assert _col(tomok_path, "간접노무비", "산업설비") == 35
    assert _col(tomok_path, "기타경비", "조경") == 41


def test_kind_col_resolved_by_header_geonchuk(geonchuk_path):
    """건축 파일은 열 배치가 달라 산업설비가 34다(하드코딩 35가 아님)."""
    assert _col(geonchuk_path, "간접노무비", "건축") == 29
    assert _col(geonchuk_path, "간접노무비", "산업설비") == 34
    assert _col(geonchuk_path, "기타경비", "산업설비") == 43


def test_missing_kind_column_raises(geonchuk_path):
    """건축 파일엔 조경 열이 없다 — 조용히 건축 값을 주지 말고 실패해야 한다."""
    with pytest.raises(LookupError):
        _col(geonchuk_path, "간접노무비", "조경")


def test_jeonmun_follows_file_base_industry(tomok_path, geonchuk_path):
    """전기통신소방전문은 파일 기본 업종(앵커 열)을 따른다.

    파일 주석: '일반관리비요율을 제외한 각종 요율은 토목, 건축 등 관련
    공사업종에 따라 적용'
    """
    assert (table_rate(tomok_path, "간접노무비", "전기통신소방전문", "10억미만", "183")
            == table_rate(tomok_path, "간접노무비", "토목", "10억미만", "183"))
    assert (table_rate(geonchuk_path, "간접노무비", "전기통신소방전문", "10억미만", "183")
            == table_rate(geonchuk_path, "간접노무비", "건축", "10억미만", "183"))
