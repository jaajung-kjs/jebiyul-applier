"""엔드투엔드: generate() 파이프라인 검증.

파이프라인: params -> compute_rates -> build_output(model→render) -> 결과 xlsx.
템플릿 복사 없이 코드로 시트를 그린다.
"""
import openpyxl
import pytest

from src.main import generate
from src.lookup import compute_rates
from src import builder


@pytest.fixture
def base_params(tomok_path):
    return dict(
        jikjeop_cost=500_000_000,
        days=200,
        kind="토목",
        contract="경쟁",
        sanjae_basis="한전",
        sanan_target=400_000_000,  # 5억미만 구간
        jebiyul_path=tomok_path,
    )


def _all_values(ws):
    return [ws.cell(r, c).value
            for r in range(1, ws.max_row + 1)
            for c in range(1, ws.max_column + 1)]


def test_generate_tomok(base_params, tmp_path):
    """generate()가 결과 xlsx를 생성하고 산재보험료(한전) 적용율을 포함한다."""
    out = generate(base_params, str(tmp_path / "out.xlsx"))
    ws = openpyxl.load_workbook(out)["적용근거"]
    assert 0.03656 in _all_values(ws)  # 산재 한전


def test_generate_ganjeop_rate_present(base_params, tmp_path):
    """간접노무비 적용율(compute_rates 참값)이 시트에 기입된다."""
    expected = compute_rates(base_params["jebiyul_path"], base_params)["간접노무비"]
    out = generate(base_params, str(tmp_path / "out2.xlsx"))
    ws = openpyxl.load_workbook(out)["적용근거"]
    vals = [v for v in _all_values(ws) if isinstance(v, (int, float))]
    assert any(abs(v - expected) < 1e-9 for v in vals)


def test_generate_does_not_raise(base_params, tmp_path):
    out = generate(base_params, str(tmp_path / "out3.xlsx"))
    assert out.endswith(".xlsx")


def test_build_output_no_template_dependency(tmp_path, tomok_path):
    """build_output이 템플릿 복사 없이 코드로 시트를 그린다(<50억 → 참조행)."""
    params = dict(kind="토목", jikjeop_cost=3_000_000_000, days=120,
                  contract="경쟁", sanjae_basis="한전", sanan_target=300_000_000)
    rates = compute_rates(tomok_path, params)
    out = str(tmp_path / "r.xlsx")
    builder.build_output(rates, out, params=params, jebiyul_path=tomok_path)
    ws = openpyxl.load_workbook(out)["적용근거"]
    assert ws["A2"].value == "공사비 산출 적용근거"
    texts = [str(v) for v in _all_values(ws) if v]
    assert any("공사규모" in t for t in texts)
    assert any("50억 이상" in t for t in texts)  # 30억 → 압축형 참조행


def test_build_output_two_sheets(tmp_path, tomok_path, hwp_path):
    from src import builder
    from src.lookup import compute_rates
    from src.nomu_model import standard_set
    params = dict(kind="토목", jikjeop_cost=500_000_000, days=200,
                  contract="경쟁", sanjae_basis="한전", sanan_target=400_000_000)
    rates = compute_rates(tomok_path, params)
    out = str(tmp_path / "two.xlsx")
    builder.build_output(rates, out, params=params, jebiyul_path=tomok_path,
                         hwp_path=hwp_path, selected_nomu=standard_set())
    wb = openpyxl.load_workbook(out)
    assert "7.통신노무임" in wb.sheetnames
    assert "8.적용근거" in wb.sheetnames
    v7 = [c.value for row in wb["7.통신노무임"].iter_rows()
          for c in row if c.value not in (None, "")]
    assert any("시중노무임 산출" in str(x) for x in v7)  # 제목
    assert any("일반공사직종" in str(x) for x in v7)  # 부문 그룹 헤더
    assert 172068 in v7  # 보통인부 현재 노임


def test_build_output_hwp_optional(tmp_path, tomok_path):
    """hwp 없이 호출하면 기존과 동일하게 단일 시트(적용근거)만 생성된다."""
    from src import builder
    from src.lookup import compute_rates
    params = dict(kind="토목", jikjeop_cost=500_000_000, days=200,
                  contract="경쟁", sanjae_basis="한전", sanan_target=400_000_000)
    rates = compute_rates(tomok_path, params)
    out = str(tmp_path / "one.xlsx")
    builder.build_output(rates, out, params=params, jebiyul_path=tomok_path)
    wb = openpyxl.load_workbook(out)
    assert wb.sheetnames == ["적용근거"]


def test_build_output_over_50_expands(tmp_path, tomok_path):
    """≥50억이면 참조행 대신 실제 규모구간이 펼쳐진다."""
    params = dict(kind="토목", jikjeop_cost=100_000_000_000, days=400,
                  contract="수의", sanjae_basis="조달청", sanan_target=20_000_000_000)
    rates = compute_rates(tomok_path, params)
    out = str(tmp_path / "big.xlsx")
    builder.build_output(rates, out, params=params, jebiyul_path=tomok_path)
    ws = openpyxl.load_workbook(out)["적용근거"]
    texts = [str(v) for v in _all_values(ws) if v]
    assert any("1000억 이상" in t for t in texts)
    assert not any("조달청 발표자료 참조" in t for t in texts)
