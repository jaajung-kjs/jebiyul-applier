"""표준 적용근거 템플릿 xlsx 생성. 율 셀은 빈 칸으로 두고 라벨만 배치한다."""
import openpyxl
from openpyxl.styles import Font

# ITEMS is the canonical list defined in the runtime package.
# Import it here so this dev-time generator stays in sync automatically.
from src.builder import ITEMS


def build(path: str = "assets/template_적용근거.xlsx") -> None:
    """빈 골격 템플릿을 생성한다.

    레이아웃 계약 (Task 8 builder와 동기화 유지):
      - A1 : 제목
      - A{3 + i*2} : 항목 라벨  (i = 0..11 → 행 3, 5, 7, … 25)
      - G{3 + i*2} : "☞ 적 용 율 :"  (설명 텍스트)
      - I{3 + i*2} : 적용율 값 — builder가 채움 (여기서는 비워 둠)
    """
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "적용근거"

    # 제목
    ws["A1"] = "공사비 산출 적용근거"
    ws["A1"].font = Font(bold=True, size=14)

    # 12개 항목 배치
    for i, item in enumerate(ITEMS):
        row = 3 + i * 2
        ws.cell(row=row, column=1, value=item).font = Font(bold=True)
        ws.cell(row=row, column=7, value="☞ 적 용 율 :")
        # column 9 (I열)은 builder가 채움 — 비워 둠

    wb.save(path)
    print(f"Template saved → {path}")


if __name__ == "__main__":
    build()
